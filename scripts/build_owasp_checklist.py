"""Build the OWASP LLM Top 10 (2026) manual checklist workbook.

The JSONL corpus under ``evals/owasp/llm-top10-2026/payloads`` is the source of
truth.  This script validates it with the standard library, renders one Excel
row per case and rewrites ``MANIFEST.sha256``.  It performs no network access.

``openpyxl`` is imported lazily and only when the workbook is rendered, so the
corpus validation, digest and ``--check`` source-side checks work without the
optional extra.

Usage:
    python scripts/build_owasp_checklist.py
    python scripts/build_owasp_checklist.py --check
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
CORPUS_ROOT = ROOT / "evals" / "owasp" / "llm-top10-2026"
PAYLOAD_ROOT = CORPUS_ROOT / "payloads"
WORKBOOK = CORPUS_ROOT / "OWASP-LLM-Top10-2026-checklist.xlsx"
MANIFEST = CORPUS_ROOT / "MANIFEST.sha256"
SCHEMA = CORPUS_ROOT / "schema.json"

PAYLOAD_FILES = (
    ("LLM01", "llm01_prompt_injection.jsonl"),
    ("LLM02", "llm02_sensitive_information.jsonl"),
    ("LLM03", "llm03_excessive_agency.jsonl"),
    ("LLM04", "llm04_supply_chain.jsonl"),
    ("LLM05", "llm05_data_model_poisoning.jsonl"),
    ("LLM06", "llm06_unbounded_consumption.jsonl"),
    ("LLM07", "llm07_misinformation.jsonl"),
    ("LLM08", "llm08_hidden_context_exposure.jsonl"),
    ("LLM09", "llm09_vector_embedding.jsonl"),
    ("LLM10", "llm10_improper_output_handling.jsonl"),
    ("CTRL", "controls_benign.jsonl"),
)

CATEGORY_NAMES = {
    "LLM01": "Prompt Injection",
    "LLM02": "Sensitive Information Disclosure",
    "LLM03": "Excessive Agency",
    "LLM04": "Supply Chain",
    "LLM05": "Data and Model Poisoning",
    "LLM06": "Unbounded Consumption",
    "LLM07": "Misinformation",
    "LLM08": "Hidden Context Exposure",
    "LLM09": "Vector and Embedding Weaknesses",
    "LLM10": "Improper Output Handling",
    "CTRL": "Benign Control",
}

OWASP_PAGES = {
    "LLM01": 10,
    "LLM02": 18,
    "LLM03": 23,
    "LLM04": 27,
    "LLM05": 33,
    "LLM06": 38,
    "LLM07": 43,
    "LLM08": 46,
    "LLM09": 50,
    "LLM10": 55,
}

REQUIRED_KEYS = (
    "id",
    "category",
    "category_name",
    "title",
    "coverage",
    "technique",
    "severity",
    "messages",
    "optional_extra",
    "guard_expect",
    "expect_status",
    "expect_code",
    "expect_behaviour",
    "control",
    "evidence",
    "needs_guard_model",
    "reference",
)

COVERAGE_VALUES = {"direct", "boundary", "known-gap", "not-applicable", "benign-control"}
SEVERITIES = {"Critical", "High", "Medium", "Low", "Info"}
GUARD_VALUES = {"should_block", "should_allow", "unverified"}
ROLES = {"user", "assistant", "system"}
STATUS_CODES = {
    200: None,
    400: "INVALID_REQUEST",
    403: "POLICY_BLOCK",
    413: "INPUT_LIMIT",
}
ALLOWED_STATUSES = set(STATUS_CODES)

RESULTS = ("Not Run", "Pass", "Fail", "Blocked", "Unverified", "N/A")
VERDICTS = ("blocked", "allowed", "error", "n/a")
FIXED_TIMESTAMP = datetime(2026, 10, 27, 0, 0, 0)

CHECKLIST_HEADERS = (
    "ID",
    "Category",
    "Category Name",
    "Coverage",
    "Technique",
    "Severity",
    "Title",
    "Payload (first user message)",
    "Messages (full JSON)",
    "Extra Request Fields",
    "Guard Expectation",
    "Expected HTTP Status",
    "Expected Error Code",
    "Expected Behaviour",
    "Control Exercised",
    "Evidence To Capture",
    "Needs Real Guard (Y/N)",
    "Reference",
    "Result",
    "Actual HTTP Status",
    "Actual Guard Verdict",
    "Latency (ms)",
    "Tester",
    "Date Tested",
    "Notes",
)

HEADER_FILL = "1F3864"
HEADER_ROW_HEIGHT = 42.0


class CorpusError(Exception):
    """The frozen corpus failed validation."""


# --------------------------------------------------------------------------- data


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise CorpusError(message)


def validate_record(record: object, *, source: str) -> dict[str, Any]:
    """Validate one record without third-party libraries."""
    _require(isinstance(record, dict), f"{source}: record must be an object")
    assert isinstance(record, dict)
    missing = [key for key in REQUIRED_KEYS if key not in record]
    _require(not missing, f"{source}: missing keys {missing}")
    unknown = set(record) - set(REQUIRED_KEYS) - {"notes"}
    _require(not unknown, f"{source}: unknown keys {sorted(unknown)}")
    _require(record["category"] in CATEGORY_NAMES, f"{source}: invalid category")
    _require(record["category_name"] == CATEGORY_NAMES[record["category"]],
             f"{source}: category_name does not match category")
    _require(isinstance(record["id"], str) and record["id"], f"{source}: invalid id")
    prefix = record["category"]
    _require(record["id"].startswith(prefix + "-"), f"{source}: id/category mismatch")
    _require(record["coverage"] in COVERAGE_VALUES, f"{source}: invalid coverage")
    _require(record["severity"] in SEVERITIES, f"{source}: invalid severity")
    _require(record["guard_expect"] in GUARD_VALUES, f"{source}: invalid guard_expect")
    _require(isinstance(record["title"], str) and record["title"].strip(), f"{source}: invalid title")
    _require(isinstance(record["technique"], str) and record["technique"].strip(),
             f"{source}: invalid technique")
    for key in ("expect_behaviour", "control", "evidence", "reference"):
        _require(isinstance(record[key], str) and record[key].strip(), f"{source}: empty {key}")
    _require(isinstance(record["needs_guard_model"], bool), f"{source}: needs_guard_model must be bool")
    _require(record["reference"].startswith("OWASP LLM") or record["reference"].startswith("Nasi Lemak"),
             f"{source}: invalid reference")
    if "notes" in record:
        _require(record["notes"] is None or isinstance(record["notes"], str), f"{source}: invalid notes")

    status = record["expect_status"]
    _require(isinstance(status, int) and not isinstance(status, bool), f"{source}: expect_status must be int")
    _require(status in ALLOWED_STATUSES, f"{source}: expect_status {status} is not an accepted gateway status")
    _require(record["expect_code"] == STATUS_CODES[status],
             f"{source}: expect_code does not match expect_status")
    _require(status != 500, f"{source}: expect_status must never be 500")

    extra = record["optional_extra"]
    _require(extra is None or isinstance(extra, dict), f"{source}: optional_extra must be object or null")

    messages = record["messages"]
    _require(isinstance(messages, list) and 1 <= len(messages) <= 16, f"{source}: invalid message count")
    for message in messages:
        _require(isinstance(message, dict), f"{source}: message must be an object")
        _require(set(message) == {"role", "content"}, f"{source}: message keys must be role/content only")
        _require(message["role"] in ROLES, f"{source}: invalid role")
        _require(isinstance(message["content"], str) and message["content"],
                 f"{source}: message content must be a non-empty string")
    return record


def load_corpus(payload_root: Path = PAYLOAD_ROOT) -> list[dict[str, Any]]:
    """Load and validate every JSONL file in the fixed order defined above."""
    records: list[dict[str, Any]] = []
    seen: set[str] = set()
    for category, filename in PAYLOAD_FILES:
        path = payload_root / filename
        if not path.is_file():
            raise CorpusError(f"missing payload file {path}")
        with path.open("r", encoding="utf-8") as handle:
            for number, line in enumerate(handle, 1):
                source = f"{filename}:{number}"
                if not line.strip():
                    raise CorpusError(f"{source}: blank line")
                try:
                    record = json.loads(line)
                except json.JSONDecodeError as exc:
                    raise CorpusError(f"{source}: invalid JSON ({exc.msg})") from exc
                validate_record(record, source=source)
                assert isinstance(record, dict)
                if record["category"] != category:
                    raise CorpusError(f"{source}: record category does not match file")
                if record["id"] in seen:
                    raise CorpusError(f"{source}: duplicate id {record['id']}")
                seen.add(record["id"])
                records.append(record)
    if len(records) < 128:
        raise CorpusError(f"corpus has {len(records)} records; at least 128 are required")
    return records


def corpus_digest(records: list[dict[str, Any]]) -> str:
    """Stable digest over the parsed corpus (independent of file whitespace)."""
    canonical = [
        json.dumps(record, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
        for record in records
    ]
    return hashlib.sha256(("\n".join(canonical) + "\n").encode("utf-8")).hexdigest()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def render_manifest(workbook_sha: str, corpus_sha: str) -> str:
    lines = [
        "# OWASP LLM Top 10 (2026) synthetic adversarial payload corpus",
        "# frozen_on: 2026-10-27",
        "# provenance: agent-authored synthetic fixtures; classifier outcomes are not established",
        f"# corpus_sha256: {corpus_sha}",
        f"{workbook_sha}  {WORKBOOK.name}",
    ]
    for _category, filename in PAYLOAD_FILES:
        relative = f"payloads/{filename}"
        lines.append(f"{file_sha256(PAYLOAD_ROOT / filename)}  {relative}")
    lines.sort(key=lambda line: (line.startswith("#"), line.split("  ", 1)[-1]))
    return "\n".join(lines) + "\n"


def manifest_entries(text: str) -> dict[str, str]:
    entries: dict[str, str] = {}
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#"):
            continue
        digest, _, relative = stripped.partition("  ")
        if len(digest) != 64 or not relative:
            raise CorpusError(f"invalid manifest line: {line!r}")
        entries[relative] = digest
    return entries


def manifest_corpus_digest(text: str) -> str | None:
    for line in text.splitlines():
        if line.startswith("# corpus_sha256: "):
            return line.split(":", 1)[1].strip()
    return None


# --------------------------------------------------------------------------- workbook


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError as exc:  # pragma: no cover - exercised only without the extra
        raise SystemExit(
            "openpyxl is required to build the workbook. Install the optional extra: "
            'pip install -e ".[eval]"'
        ) from exc
    return openpyxl


def _first_user_message(record: dict[str, Any]) -> str:
    for message in record["messages"]:
        if message["role"] == "user":
            return message["content"]
    return record["messages"][0]["content"]


_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _cell_text(value: str) -> str:
    """Render control characters as escapes; openpyxl cannot store them in XML.

    The JSONL file keeps the exact bytes.  Only the workbook display is escaped.
    """
    return _ILLEGAL_XML.sub(lambda match: "\\u%04x" % ord(match.group()), value)


def _extra_display(record: dict[str, Any]) -> str:
    extra = record["optional_extra"]
    if extra is None:
        return ""
    return json.dumps(extra, ensure_ascii=False, sort_keys=True)


def _style_header(ws, row: int, columns: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type="solid")
    font = Font(bold=True, color="FFFFFF")
    for column in range(1, columns + 1):
        cell = ws.cell(row=row, column=column)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = HEADER_ROW_HEIGHT


def _sheet_read_me(wb, records: list[dict[str, Any]], corpus_sha: str) -> None:
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("Read Me")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 110
    rows = [
        ("OWASP LLM Top 10 (2026) manual checklist", ""),
        ("Purpose", "A frozen, synthetic adversarial payload corpus for the Nasi Lemak gateway, rendered as a hand-run checklist. Each row carries the exact request, the expected gateway result and the evidence to capture."),
        ("Scope", "Black-box, manual testing of the HTTP gateway and its guard wiring. No runner, no local model weights and no network probes are part of this slice."),
        ("OWASP source", "OWASP Top 10 for LLM Applications 2026, v1.0 (August 2026). Page references in each row point to the section start in that PDF."),
        ("Coverage honesty", "A pass on a 'boundary' or 'known-gap' row is not detection coverage: those rows record that the surface does not exist, or that the gateway is expected to miss. Use the Summary sheet counts, not the raw pass rate."),
        ("Classifier quality", "The guard models are reached over an API. Their precision and recall are unproven until real artifacts and thresholds are recorded in the Run Log; no row here establishes classifier quality."),
        ("How to run a session", "1) Confirm the endpoint, policy version and guard model/threshold in Run Log. 2) Work top to bottom. 3) Send the request body implied by the messages plus the extra fields (fixture-only fields starting with '$' are never sent). 4) Record the actual status, guard verdict and latency. 5) Set Result to Pass, Fail, Blocked, Unverified or N/A with a note."),
        ("Result legend", "Not Run = untouched. Pass = observed behaviour matches Expected Behaviour. Fail = it does not. Blocked = the request could not be sent (environment, auth, endpoint). Unverified = the gateway answered but the expected property cannot be checked manually (for example truthfulness). N/A = the surface is absent in this build."),
        ("Fixture-only fields", "optional_extra keys beginning with '$' (such as $fixture_reply and $raw_body) are tester setup notes, not request fields. Everything else is merged into the request body and may make it invalid on purpose."),
        ("Corpus digest", corpus_sha),
        ("Case count", str(len(records))),
        ("Regenerate", r'.\.venv\Scripts\python.exe scripts/build_owasp_checklist.py'),
        ("Verify", r'.\.venv\Scripts\python.exe -m unittest tests.test_owasp_suite -v'),
    ]
    for row_index, (label, value) in enumerate(rows, 1):
        ws.cell(row=row_index, column=1, value=label)
        ws.cell(row=row_index, column=2, value=value)
        ws.cell(row=row_index, column=1).font = Font(bold=True)
        ws.cell(row=row_index, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=row_index, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[row_index].height = 18 if not value else max(18, 15 * (1 + len(value) // 105))
    ws.freeze_panes = "A2"


def _sheet_checklist(wb, records: list[dict[str, Any]]) -> None:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet("Checklist")
    ws.append(list(CHECKLIST_HEADERS))
    _style_header(ws, 1, len(CHECKLIST_HEADERS))

    for record in records:
        ws.append([
            _cell_text(record["id"]),
            _cell_text(record["category"]),
            _cell_text(record["category_name"]),
            _cell_text(record["coverage"]),
            _cell_text(record["technique"]),
            _cell_text(record["severity"]),
            _cell_text(record["title"]),
            _cell_text(_first_user_message(record)),
            _cell_text(json.dumps(record["messages"], ensure_ascii=False)),
            _cell_text(_extra_display(record)),
            _cell_text(record["guard_expect"]),
            record["expect_status"],
            _cell_text(record["expect_code"]) if record["expect_code"] else None,
            _cell_text(record["expect_behaviour"]),
            _cell_text(record["control"]),
            _cell_text(record["evidence"]),
            "Y" if record["needs_guard_model"] else "N",
            _cell_text(record["reference"]),
            "Not Run",
            None,
            None,
            None,
            None,
            None,
            _cell_text(record.get("notes") or "") or None,
        ])

    last = len(records) + 1
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_columns = (7, 8, 9, 10, 14, 15, 16, 25)
    for row in ws.iter_rows(min_row=2, max_row=last, max_col=len(CHECKLIST_HEADERS)):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in wrap_columns)
    for cell in ws[1]:
        cell.border = border

    widths = {
        "A": 11, "B": 8, "C": 30, "D": 14, "E": 30, "F": 10, "G": 40,
        "H": 52, "I": 64, "J": 40, "K": 16, "L": 10, "M": 18, "N": 58,
        "O": 30, "P": 44, "Q": 12, "R": 24, "S": 12, "T": 10, "U": 16,
        "V": 10, "W": 14, "X": 12, "Y": 42,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:Y{last}"

    result_validation = DataValidation(
        type="list", formula1='"' + ",".join(RESULTS) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid result",
        error="Choose Not Run, Pass, Fail, Blocked, Unverified or N/A.",
    )
    result_validation.add(f"S2:S{last}")
    ws.add_data_validation(result_validation)

    verdict_validation = DataValidation(
        type="list", formula1='"' + ",".join(VERDICTS) + '"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid verdict",
        error="Choose blocked, allowed, error or n/a.",
    )
    verdict_validation.add(f"U2:U{last}")
    ws.add_data_validation(verdict_validation)

    def fill(colour: str) -> PatternFill:
        return PatternFill(start_color=colour, end_color=colour, fill_type="solid")

    for value, colour in (
        ("Pass", "C6EFCE"),
        ("Fail", "FFC7CE"),
        ("Blocked", "FFEB9C"),
        ("Unverified", "D9D9D9"),
        ("N/A", "EDEDED"),
    ):
        ws.conditional_formatting.add(
            f"S2:S{last}",
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill(colour)),
        )

    ws.print_title_rows = "1:1"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _sheet_summary(wb, records: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet("Summary")
    last = len(records) + 1
    ws["A1"] = "Coverage and results summary"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Formulas read the Checklist sheet; a high pass rate on boundary or known-gap rows is not detection coverage."
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 30

    headers = ["Category", "Name", "Cases", "Pass", "Fail", "Blocked", "Unverified", "N/A", "Not Run", "Pass rate (Pass/(Pass+Fail))"]
    ws.append([])
    ws.append(headers)
    _style_header(ws, 4, len(headers))
    categories = [category for category, _filename in PAYLOAD_FILES]
    for offset, category in enumerate(categories):
        row = 5 + offset
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=CATEGORY_NAMES[category])
        ws.cell(row=row, column=3, value=f'=COUNTIF(Checklist!$B$2:$B${last},$A{row})')
        for column, result in ((4, "Pass"), (5, "Fail"), (6, "Blocked"), (7, "Unverified"), (8, "N/A"), (9, "Not Run")):
            ws.cell(row=row, column=column,
                    value=f'=COUNTIFS(Checklist!$B$2:$B${last},$A{row},Checklist!$S$2:$S${last},"{result}")')
        ws.cell(row=row, column=10, value=f'=IFERROR(D{row}/(D{row}+E{row}),"")')
        ws.cell(row=row, column=10).number_format = "0.0%"

    total_row = 5 + len(categories)
    ws.cell(row=total_row, column=1, value="Total")
    ws.cell(row=total_row, column=2, value="All cases")
    for column in range(3, 10):
        letter = chr(ord("A") + column - 1)
        ws.cell(row=total_row, column=column, value=f"=SUM({letter}5:{letter}{total_row - 1})")
    ws.cell(row=total_row, column=10, value=f'=IFERROR(D{total_row}/(D{total_row}+E{total_row}),"")')
    ws.cell(row=total_row, column=10).number_format = "0.0%"
    for cell in ws[total_row]:
        cell.font = Font(bold=True)

    coverage_row = total_row + 2
    ws.cell(row=coverage_row, column=1, value="Coverage caveats").font = Font(bold=True, size=12)
    coverage_rows = [
        ("Boundary cases (surface absent)", f'=COUNTIF(Checklist!$D$2:$D${last},"boundary")',
         "The feature does not exist in this build; a clean refusal is not a detection."),
        ("Known-gap cases (expected to fail)", f'=COUNTIF(Checklist!$D$2:$D${last},"known-gap")',
         "Documented gaps; do not count an expected failure as a pass."),
        ("Not-applicable cases", f'=COUNTIF(Checklist!$D$2:$D${last},"not-applicable")',
         "No ingest, RAG, vector or training surface exists."),
        ("Benign controls", f'=COUNTIF(Checklist!$D$2:$D${last},"benign-control")',
         "Legitimate traffic; a block here is a false positive."),
        ("Cases needing the real guard model", f'=COUNTIF(Checklist!$Q$2:$Q${last},"Y")',
         "Outcome depends on the external classifier; unproven until real artifacts are used."),
        ("Cases independent of the guard model", f'=COUNTIF(Checklist!$Q$2:$Q${last},"N")',
         "Deterministic schema, limit, masking or buffering checks."),
    ]
    for offset, (label, formula, note) in enumerate(coverage_rows, coverage_row + 1):
        ws.cell(row=offset, column=1, value=label)
        ws.cell(row=offset, column=3, value=formula)
        ws.cell(row=offset, column=4, value=note)
        ws.merge_cells(start_row=offset, start_column=4, end_row=offset, end_column=10)

    results_row = coverage_row + len(coverage_rows) + 2
    ws.cell(row=results_row, column=1, value="Results by value").font = Font(bold=True, size=12)
    for offset, result in enumerate(RESULTS, results_row + 1):
        ws.cell(row=offset, column=1, value=result)
        ws.cell(row=offset, column=3, value=f'=COUNTIF(Checklist!$S$2:$S${last},"{result}")')

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 34
    for column in "CDEFGHIJ":
        ws.column_dimensions[column].width = 12


def _sheet_category_map(wb, records: list[dict[str, Any]]) -> None:
    from openpyxl.styles import Alignment

    ws = wb.create_sheet("Category Map")
    headers = ["ID", "Name", "What it means", "Testable in this build", "Primary control", "OWASP page", "Cases"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    last = len(records) + 1
    meanings = {
        "LLM01": "Hostile input tries to override or escape the model's instructions, directly or through quoted content.",
        "LLM02": "The system exposes or fails to mask confidential, regulated or personal data.",
        "LLM03": "Granted tools, functions or autonomy let model output reach privileged actions.",
        "LLM04": "Compromised models, adapters, packages or providers enter the model supply chain.",
        "LLM05": "Adversarial or unsafe data and labels corrupt training, fine-tuning or retrieval content.",
        "LLM06": "Excessive or uncontrolled inference, tokens, body size or concurrency consumes resources.",
        "LLM07": "Confident but incorrect or unsupported output drives a harmful decision.",
        "LLM08": "Hidden system prompts, rules, tools or configuration are extracted or reconstructed.",
        "LLM09": "Retrieval and embedding geometry leaks data or lets poisoned chunks dominate results.",
        "LLM10": "Model output reaches a downstream HTML, CSV, shell or template sink without escaping.",
        "CTRL": "Legitimate traffic used to measure false positives on the same path.",
    }
    testable = {
        "LLM01": "Input attack guard and literal deny phrases across every accepted turn.",
        "LLM02": "Deterministic input and output PII masking for EMAIL, MY_PHONE and MYKAD.",
        "LLM03": "Schema rejection: no tools, functions or agent surface exists.",
        "LLM04": "Pinned model identity; unknown fields are rejected and downloads never happen.",
        "LLM05": "Not testable: no ingest, fine-tuning or persistent memory path.",
        "LLM06": "Message count, 4096-character serialization, 64 KiB body, max_tokens and stream.",
        "LLM07": "Not testable as truth: the gateway screens safety only and does not verify facts.",
        "LLM08": "No hidden context exists; checks that nothing privileged is fabricated or leaked.",
        "LLM09": "Not testable: no vector store, embedding model, RAG or tenant corpus.",
        "LLM10": "Known gap: buffered JSON only; the consumer must escape output.",
        "CTRL": "False-positive rate on benign prompts, including security education and PII in context.",
    }
    controls = {
        "LLM01": "input_attack_guard / input_deny_phrase",
        "LLM02": "pii_input_masking / pii_output_masking",
        "LLM03": "request_schema_validation",
        "LLM04": "request_schema_validation",
        "LLM05": "no-ingest-surface",
        "LLM06": "body_size_limit / input_token_budget / request_schema_validation",
        "LLM07": "output_guard/safety-only",
        "LLM08": "output_guard/no-hidden-context",
        "LLM09": "no-vector-surface",
        "LLM10": "output_handling/consumer-escaping",
        "CTRL": "false-positive-control",
    }
    llm_categories = [(category, filename) for category, filename in PAYLOAD_FILES if category != "CTRL"]
    for row, (category, _filename) in enumerate(llm_categories, 2):
        page = OWASP_PAGES.get(category)
        ws.cell(row=row, column=1, value=category)
        ws.cell(row=row, column=2, value=CATEGORY_NAMES[category])
        ws.cell(row=row, column=3, value=meanings[category])
        ws.cell(row=row, column=4, value=testable[category])
        ws.cell(row=row, column=5, value=controls[category])
        ws.cell(row=row, column=6, value=page if page else "n/a")
        ws.cell(row=row, column=7, value=f'=COUNTIF(Checklist!$B$2:$B${last},$A{row})')
        for cell in ws[row]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 10, "B": 32, "C": 60, "D": 55, "E": 38, "F": 12, "G": 8}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:G{len(llm_categories) + 1}"


def _sheet_gaps(wb) -> None:
    from openpyxl.styles import Alignment

    ws = wb.create_sheet("Known Gaps")
    headers = ["ID", "Known gap", "Impact on this checklist", "Where documented", "Status"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    gaps = [
        ("KG-01", "MyKad recogniser false positives (shape-only, invalid dates are masked).",
         "LLM02-009 passes masking but should be recorded as a false positive, not privacy coverage.", "docs/pii-quality.md", "Open"),
        ("KG-02", "Passport numbers are unsupported.", "LLM02-010 is expected to leak; do not count it as a pass.", "docs/pii-quality.md", "Open"),
        ("KG-03", "Bank account numbers are unsupported.", "LLM02-011 is expected to leak; do not count it as a pass.", "docs/pii-quality.md", "Open"),
        ("KG-04", "Spaced international phone numbers such as '+60 12-345 6789' are not matched.", "LLM02-004 is a known false negative; the compact 60 form is the positive control.", "docs/pii-quality.md", "Open"),
        ("KG-05", "PII split across messages or turns is not stitched.", "LLM02-013 is expected to leak fragments.", "docs/pii-quality.md", "Open"),
        ("KG-06", "No RAG, vector store, embedding model or tenant corpus exists.", "LLM09 rows are not-applicable; retrieval behaviour cannot be demonstrated here.", "docs/pilot-scope.md", "By design"),
        ("KG-07", "No tools, functions, agents or MCP surface exists.", "LLM03 rows test schema rejection only; no execution path is exercised.", "docs/api-contract.yaml", "By design"),
        ("KG-08", "No ingest, fine-tuning or persistent memory path exists.", "LLM05 rows are not-applicable and cannot demonstrate poisoning coverage.", "docs/data-handling.md", "By design"),
        ("KG-09", "Classifier quality is unverified without real guard artifacts and thresholds.", "Every needs_guard_model=Y row is only as good as the deployed classifier; record the model and threshold.", "docs/guards.md", "Blocked on artifacts"),
        ("KG-10", "The gateway does not verify truth.", "LLM07 rows record unverified claims; a 200 is not an accuracy pass.", "docs/threat-model.md", "By design"),
        ("KG-11", "The gateway does not sanitise output for downstream HTML, CSV, ANSI or templates.", "LLM10 rows are known-gap: buffered JSON is returned; the consumer must escape.", "docs/api-contract.yaml", "Open"),
        ("KG-12", "A high pass rate on boundary or known-gap rows is not detection coverage.", "Read Summary caveat counts alongside the pass rate.", "docs/owasp-payloads.md", "Reporting rule"),
    ]
    for row, gap in enumerate(gaps, 2):
        for column, value in enumerate(gap, 1):
            cell = ws.cell(row=row, column=column, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 8, "B": 62, "C": 70, "D": 28, "E": 18}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{len(gaps) + 1}"


def _sheet_run_log(wb) -> None:
    from openpyxl.styles import Alignment, Border, Side

    ws = wb.create_sheet("Run Log")
    headers = ["Date", "Tester", "Build / commit", "Endpoint", "Guard model and threshold",
               "Cases run", "Verdict", "Notes"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in range(2, 22):
        for column in range(1, len(headers) + 1):
            cell = ws.cell(row=row, column=column)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for column, width in {"A": 14, "B": 18, "C": 24, "D": 34, "E": 38, "F": 10, "G": 14, "H": 60}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"


def build_workbook_bytes(records: list[dict[str, Any]]) -> bytes:
    """Render the workbook in memory and return deterministic bytes."""
    openpyxl = _openpyxl()
    from openpyxl import Workbook

    if len(records) < 1:
        raise CorpusError("cannot build a workbook from an empty corpus")
    corpus_sha = corpus_digest(records)
    wb = Workbook()
    wb.remove(wb.active)

    _sheet_read_me(wb, records, corpus_sha)
    _sheet_checklist(wb, records)
    _sheet_summary(wb, records)
    _sheet_category_map(wb, records)
    _sheet_gaps(wb)
    _sheet_run_log(wb)

    wb.properties.creator = "Nasi Lemak evaluator"
    wb.properties.lastModifiedBy = "Nasi Lemak evaluator"
    wb.properties.title = "OWASP LLM Top 10 (2026) checklist"
    wb.properties.subject = "Synthetic adversarial payload corpus"
    wb.properties.description = "Manual black-box checklist; corpus digest " + corpus_sha
    wb.properties.keywords = "OWASP,LLM,2026,synthetic,guardrails"
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP
    wb.properties.lastPrinted = FIXED_TIMESTAMP

    buffer = io.BytesIO()
    wb.save(buffer)
    return _normalise_zip(buffer.getvalue())


def _normalise_zip(data: bytes) -> bytes:
    """Rewrite the xlsx zip with fixed metadata so rebuilds are byte-identical."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl stamps the current time into dcterms:modified on every
                # save; pin it so a rebuild is byte-identical.
                payload = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                    rb"\g<1>2026-10-27T00:00:00Z\g<2>",
                    payload,
                )
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            fixed.create_system = 0
            target.writestr(fixed, payload)
    return output.getvalue()


def build(payload_root: Path = PAYLOAD_ROOT, workbook: Path = WORKBOOK,
          manifest: Path = MANIFEST) -> dict[str, Any]:
    records = load_corpus(payload_root)
    data = build_workbook_bytes(records)
    if len(data) > 2 * 1024 * 1024:
        raise CorpusError(f"workbook is {len(data)} bytes; keep it under 2 MiB")
    workbook.write_bytes(data)
    workbook_sha = hashlib.sha256(data).hexdigest()
    manifest.write_text(render_manifest(workbook_sha, corpus_digest(records)), encoding="utf-8", newline="\n")
    return {
        "records": len(records),
        "workbook_bytes": len(data),
        "workbook_sha256": workbook_sha,
        "corpus_sha256": corpus_digest(records),
    }


def check(payload_root: Path = PAYLOAD_ROOT, workbook: Path = WORKBOOK,
          manifest: Path = MANIFEST) -> list[str]:
    """Return drift problems without writing anything."""
    problems: list[str] = []
    records = load_corpus(payload_root)
    if not workbook.is_file():
        problems.append("workbook is missing")
    if not manifest.is_file():
        problems.append("manifest is missing")
        return problems
    text = manifest.read_text(encoding="utf-8")
    entries = manifest_entries(text)
    if entries.get(WORKBOOK.name) != file_sha256(workbook):
        problems.append("workbook hash does not match MANIFEST.sha256")
    for _category, filename in PAYLOAD_FILES:
        relative = f"payloads/{filename}"
        if entries.get(relative) != file_sha256(payload_root / filename):
            problems.append(f"{relative} does not match MANIFEST.sha256")
    if manifest_corpus_digest(text) != corpus_digest(records):
        problems.append("corpus digest does not match MANIFEST.sha256")
    openpyxl_available = True
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        openpyxl_available = False
    if openpyxl_available:
        if build_workbook_bytes(records) != workbook.read_bytes():
            problems.append("workbook drifted from the corpus; regenerate it")
    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true",
                        help="verify the committed workbook and manifest without writing")
    args = parser.parse_args(argv)
    try:
        if args.check:
            problems = check()
            if problems:
                for problem in problems:
                    print(f"FAIL: {problem}", file=sys.stderr)
                return 1
            print("ok: corpus, workbook and manifest agree")
            return 0
        summary = build()
    except CorpusError as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    print(f"records={summary['records']} workbook_bytes={summary['workbook_bytes']}")
    print(f"workbook_sha256={summary['workbook_sha256']}")
    print(f"corpus_sha256={summary['corpus_sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
